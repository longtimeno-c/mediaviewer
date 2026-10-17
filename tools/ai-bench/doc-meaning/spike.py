#!/usr/bin/env python3
# Copyright (C) 2026 longtimeno-c
# SPDX-License-Identifier: GPL-3.0-or-later
"""#354 slice D3's model spike (docs/plans/document-search.md §1): which small
sentence-embedding model ranks document passages best for natural-language
queries, and what it costs.

  spike.py passages <corpus> <work>        text -> passages (PDFium and python-docx)
  spike.py embed <work> <models> <model> <precision> [threads]
  spike.py speed <work> <models> <model> <precision> [threads=2] [n=1024]
  spike.py score <work> <queries.json> <corpus> [<model>/<precision> ...]
  spike.py golden <models> <model> <precision> <out.json>   ORT-Python reference vectors

A document scores as its best passage (the engine merges passages per document,
best page first). P@5 is relevant hits in the top 5 over min(5, relevant);
R@10 is relevant hits in the top 10 over min(10, relevant). Run with -I.
"""
import json
import math
import os
import re
import sys
import time

import numpy as np

MODELS = {
    # name: (pooling, query prefix, passage prefix)
    "minilm": ("mean", "", ""),
    "bge": ("cls", "Represent this sentence for searching relevant passages: ", ""),
    "arctic": ("cls", "Represent this sentence for searching relevant passages: ", ""),
    "e5": ("mean", "query: ", "passage: "),
}
FILES = {"fp32": "onnx/model.onnx", "fp16": "onnx/model_fp16.onnx", "int8": "onnx/model_int8.onnx",
         "q8": "onnx/model_quantized.onnx"}
MAX_TOKENS = 256
MAX_PAGES = 500
# Passages per document the spike embeds (its first ones): a 300-page IRS
# publication is thousands, and four models at two precisions over all of
# them took hours on a shared machine. The engine embeds every passage.
CAP = int(os.environ.get("MV_SPIKE_CAP", "64"))

# ---- passages: the rule of src/addons/ai/passages.cpp --------------------------------------

PLAIN = str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-",
                       "—": "-", " ": " ", "ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff",
                       "ﬃ": "ffi", "ﬄ": "ffl", "­": ""})


def pdf_page_prose(t):
    t = t.replace("\r\n", "\n").translate(PLAIN)
    t = re.sub(r"(\w)-\n(\w)", r"\1\2", t)
    return " ".join(t.split())


def build_passages(blocks):
    """blocks: (page, text, block_end). Returns [(page, text)]."""
    out, cur, page, words = [], [], 0, 0

    def close():
        nonlocal cur, words
        if cur:
            out.append((page, " ".join(cur)))
        cur, words = [], 0

    for pg, text, block_end in blocks:
        if pg != page:
            close()
            page = pg
        for w in text.split():
            cur.append(w)
            words += 1
            if (words >= 40 and re.search(r"[.!?][\"')\]]*$", w)) or words >= 80:
                close()
        if block_end and words >= 40:
            close()
    close()
    return out


def extract(path):
    blocks = []
    if path.endswith(".pdf"):
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(path)
        for i in range(min(len(doc), MAX_PAGES)):
            tp = doc[i].get_textpage()
            blocks.append((i, pdf_page_prose(tp.get_text_range()), True))
    else:
        import docx
        d = docx.Document(path)
        for p in d.paragraphs:
            t = " ".join(p.text.translate(PLAIN).split())
            if t:
                blocks.append((0, t, True))
    return build_passages(blocks)


def cmd_passages(corpus, work):
    os.makedirs(work, exist_ok=True)
    names = sorted(json.load(open(os.path.join(corpus, "corpus.json"))))
    rows = []
    for n in names:
        try:
            ps = extract(os.path.join(corpus, "docs", n))
        except Exception as e:  # noqa: BLE001
            print("  unreadable", n, e, file=sys.stderr)
            continue
        rows += [{"doc": n, "page": p, "text": t} for p, t in ps[:CAP]]
    json.dump(rows, open(os.path.join(work, "passages.json"), "w"), ensure_ascii=False)
    print(len(names), "documents", len(rows), "passages")

# ---- embedding ------------------------------------------------------------------------------


class Embedder:
    def __init__(self, models, name, precision, threads):
        import onnxruntime as ort
        from tokenizers import Tokenizer
        folder = os.path.join(models, name)
        so = ort.SessionOptions()
        so.intra_op_num_threads = threads
        so.inter_op_num_threads = 1
        so.log_severity_level = 3
        self.sess = ort.InferenceSession(os.path.join(folder, FILES[precision]), so,
                                         providers=["CPUExecutionProvider"])
        self.inputs = {i.name for i in self.sess.get_inputs()}
        self.tok = Tokenizer.from_file(os.path.join(folder, "tokenizer.json"))
        # Xenova's tokenizer.json pads to a fixed 128: the batch pads itself.
        self.tok.no_padding()
        self.tok.enable_truncation(MAX_TOKENS)
        self.pool, self.qpre, self.ppre = MODELS[name]

    def embed(self, texts, batch=16):
        out = []
        for s in range(0, len(texts), batch):
            enc = self.tok.encode_batch(texts[s:s + batch])
            n = max(len(e.ids) for e in enc)
            ids = np.zeros((len(enc), n), np.int64)
            mask = np.zeros((len(enc), n), np.int64)
            for i, e in enumerate(enc):
                ids[i, :len(e.ids)] = e.ids
                mask[i, :len(e.ids)] = 1
            feed = {"input_ids": ids, "attention_mask": mask}
            if "token_type_ids" in self.inputs:
                feed["token_type_ids"] = np.zeros_like(ids)
            h = self.sess.run(None, feed)[0].astype(np.float32)
            if self.pool == "cls":
                v = h[:, 0]
            else:
                m = mask[..., None].astype(np.float32)
                v = (h * m).sum(1) / np.maximum(m.sum(1), 1e-9)
            v /= np.linalg.norm(v, axis=1, keepdims=True)
            out.append(v)
        return np.concatenate(out)

    def passages(self, texts, batch=16):
        return self.embed([self.ppre + t for t in texts], batch)

    def queries(self, texts):
        return self.embed([self.qpre + t for t in texts], 1)


def cmd_embed(work, models, name, precision, threads=8):
    rows = json.load(open(os.path.join(work, "passages.json")))
    e = Embedder(models, name, precision, int(threads))
    # Length-sorted batches waste no padding; the order is restored after.
    order = sorted(range(len(rows)), key=lambda i: len(rows[i]["text"]))
    t0 = time.time()
    v = e.passages([rows[i]["text"] for i in order], batch=32)
    out = np.empty_like(v)
    out[order] = v
    np.save(os.path.join(work, f"emb-{name}-{precision}.npy"), out.astype(np.float16))
    print(name, precision, len(rows), "passages", f"{time.time() - t0:.0f} s")


def cmd_speed(work, models, name, precision, threads=2, n=1024):
    rows = json.load(open(os.path.join(work, "passages.json")))
    step = max(1, len(rows) // int(n))
    texts = [r["text"] for r in rows[::step]][:int(n)]
    e = Embedder(models, name, precision, int(threads))
    e.passages(texts[:32])
    t0 = time.time()
    e.passages(texts, batch=16)
    dt = time.time() - t0
    q = ["a query of about eight words long here"] * 50
    t1 = time.time()
    e.queries(q)
    qdt = (time.time() - t1) / len(q)
    print(json.dumps({"model": name, "precision": precision, "threads": int(threads), "passages": len(texts),
                      "passages_per_s": round(len(texts) / dt, 1), "query_ms": round(qdt * 1000, 2)}))


def cmd_golden(models, name, precision, out):
    texts = ["A receipt for two coffees and a croissant, paid by card.",
             "The tenant shall pay the rent on the first day of each month.",
             "Graph neural networks pass messages along the edges of a graph.",
             "Le Vésuve est un volcan situé au bord de la baie de Naples.",
             "Ünïcödé — “quotes” and ligatures ﬁnd their way in.",
             ""]
    e = Embedder(models, name, precision, 1)
    p = e.passages(texts)
    q = e.queries(["tax return from my accountant"])
    json.dump({"model": name, "precision": precision, "passages": texts,
               "passage_vecs": p.round(6).tolist(), "query": "tax return from my accountant",
               "query_vec": q[0].round(6).tolist()}, open(out, "w"), ensure_ascii=False)

# ---- scoring --------------------------------------------------------------------------------


def stem(n):
    return n.rsplit(".", 1)[0]


def relevant_set(q, corpus):
    rel = set()
    stems = {}
    for n in corpus:
        stems.setdefault(stem(n), []).append(n)
    for r in q["relevant"]:
        if r.startswith("subject:"):
            rel |= {n for n, m in corpus.items() if m["subject"] == r[8:]}
        else:
            rel |= set(stems.get(stem(r), []))
    return rel


def bm25_doc_scores(rows, queries):
    from collections import Counter
    tok = re.compile(r"[a-z0-9]+")
    lens, post = [], {}
    for i, r in enumerate(rows):
        words = tok.findall(r["text"].lower())
        lens.append(len(words))
        for w, tf in Counter(words).items():
            post.setdefault(w, []).append((i, tf))
    N, avg = len(rows), sum(lens) / len(rows)
    stop = set("a an the of in on to for and or my i me is it at by with from that this how what".split())
    out = []
    for q in queries:
        s = np.zeros(N)
        for w in dict.fromkeys(t for t in tok.findall(q.lower()) if t not in stop):
            p = post.get(w, [])
            if not p:
                continue
            idf = math.log(1 + (N - len(p) + 0.5) / (len(p) + 0.5))
            for i, tf in p:
                s[i] += idf * tf * 2.2 / (tf + 1.2 * (0.25 + 0.75 * lens[i] / avg))
        out.append(s)
    return np.array(out)


def rank_docs(scores, doc_of, k=10):
    best = {}
    for i in np.argsort(-scores)[:5000]:
        d = doc_of[i]
        if d not in best:
            best[d] = scores[i]
            if len(best) >= k:
                break
    return list(best)


def cmd_score(work, queries_path, corpus_dir, *configs):
    rows = json.load(open(os.path.join(work, "passages.json")))
    corpus = json.load(open(os.path.join(corpus_dir, "corpus.json")))
    indexed = {r["doc"] for r in rows}
    corpus = {n: m for n, m in corpus.items() if n in indexed}
    qs = json.load(open(queries_path))["queries"]
    doc_of = [r["doc"] for r in rows]
    rels = [relevant_set(q, corpus) for q in qs]
    for q, r in zip(qs, rels):
        if not r:
            print("  no relevant document indexed for:", q["q"], file=sys.stderr)
    configs = list(configs) or sorted(f[4:-4].replace("-", "/", 1) for f in os.listdir(work) if f.startswith("emb-"))
    results = {}

    def evaluate(label, score_rows):
        per = []
        for q, rel, s in zip(qs, rels, score_rows):
            top = rank_docs(s, doc_of)
            p5 = len(set(top[:5]) & rel) / max(1, min(5, len(rel)))
            r10 = len(set(top[:10]) & rel) / max(1, min(10, len(rel)))
            per.append({"q": q["q"], "p5": round(p5, 3), "r10": round(r10, 3), "top5": top[:5]})
        sets = {"all": range(len(qs)),
                "english": [i for i, q in enumerate(qs) if q.get("lang", "en") == "en"],
                "other_language": [i for i, q in enumerate(qs) if q.get("lang", "en") != "en"]}
        summary = {k: {"p5": round(float(np.mean([per[i]["p5"] for i in v])), 3),
                       "r10": round(float(np.mean([per[i]["r10"] for i in v])), 3), "n": len(v)}
                   for k, v in sets.items() if len(v)}
        results[label] = {"summary": summary, "per_query": per}
        print(label.ljust(16), " ".join(f"{k} P@5 {v['p5']:.3f} R@10 {v['r10']:.3f} (n={v['n']})"
                                        for k, v in summary.items()))

    evaluate("bm25", bm25_doc_scores(rows, [q["q"] for q in qs]))
    for c in configs:
        name, precision = c.split("/")
        emb = np.load(os.path.join(work, f"emb-{name}-{precision}.npy")).astype(np.float32)
        e = Embedder(os.environ.get("MV_SPIKE_MODELS", os.path.join(work, "..", "models")), name, precision, 4)
        qv = e.queries([q["q"] for q in qs])
        evaluate(c, qv @ emb.T)
    json.dump(results, open(os.path.join(work, "scores.json"), "w"), indent=1, ensure_ascii=False)


GENERIC = ["a document.", "some text.", "a page of a document.", "this is a paragraph of text.", "information."]


def cmd_calibrate(work, queries_path, models, config):
    """The 'nothing found' margin (the CLAP / PR 20 rule): a passage's margin is its
    cosine minus its best cosine against the generic prompts (embedded as queries).
    A query answers when a top-ten passage beats that by query_margin."""
    name, precision = config.split("/")
    emb = np.load(os.path.join(work, f"emb-{name}-{precision}.npy")).astype(np.float32)
    e = Embedder(models, name, precision, 4)
    g = e.queries(GENERIC)
    generic = (emb @ g.T).max(1)
    d = json.load(open(queries_path))
    real = [q["q"] for q in d["queries"]]
    none = d["nothing"]

    def best(qs):
        v = e.queries(qs) @ emb.T
        out = []
        for s in v:
            top = np.argsort(-s)[:10]
            out.append(float((s[top] - generic[top]).max()))
        return np.array(out)

    br, bn = best(real), best(none)
    qs = lambda a: " ".join(f"{x:.3f}" for x in np.quantile(a, [0, 0.1, 0.25, 0.5, 0.75, 0.9, 1]))
    print(config, "real  min/10/25/50/75/90/max:", qs(br))
    print(config, "none  min/10/25/50/75/90/max:", qs(bn))
    best_t, best_err = 0, 1e9
    for t in np.arange(0.0, 1.0, 0.0025):
        err = (br < t).sum() + (bn >= t).sum()
        if err < best_err:
            best_t, best_err = t, err
    print(config, f"query_margin {best_t:.4f}: real answered {(br >= best_t).sum()}/{len(br)}, "
                  f"nothing answered {(bn >= best_t).sum()}/{len(bn)}")
    print(config, "generic mean", float(generic.mean()))
    json.dump({"config": config, "generic_prompts": GENERIC, "query_margin": round(float(best_t), 4),
               "real_best": br.round(4).tolist(), "nothing_best": bn.round(4).tolist(), "nothing": none},
              open(os.path.join(work, f"calibrate-{name}-{precision}.json"), "w"), indent=1)


def main():
    cmd, args = sys.argv[1], sys.argv[2:]
    {"passages": cmd_passages, "embed": cmd_embed, "speed": cmd_speed, "score": cmd_score,
     "golden": cmd_golden, "calibrate": cmd_calibrate}[cmd](*args)


if __name__ == "__main__":
    main()
