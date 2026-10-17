#!/usr/bin/env python3
# Copyright (C) 2026 longtimeno-c
# SPDX-License-Identifier: GPL-3.0-or-later
"""#354 slice D3's spike corpus (docs/plans/document-search.md §1): about 500 public
documents with known subjects, so 50 natural-language queries can be judged by
what each document is about.

  - IRS forms, instructions and publications (real PDFs, irs.gov)
  - arXiv papers found by title words on 15 subjects (real PDFs, export.arxiv.org)
  - Wikipedia articles on 12 subjects in English, plus some in French, German and
    Spanish (real text; written here as DOCX or a plain PDF, since public
    DOCX in bulk does not exist)

Usage: fetch_corpus.py <out>     (writes <out>/docs/* and <out>/corpus.json)
Run it with -I: the files it fetches are untrusted input.
"""
import json
import os
import re
import sys
import time
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET

UA = "MediaViewer-doc-search-spike/1 (local benchmark; contact via github.com/longtimeno-c)"

IRS = [
    # forms
    "f1040", "f1040s1", "f1040s2", "f1040s3", "f1040sa", "f1040sb", "f1040sc", "f1040sd", "f1040se",
    "f1040sse", "f1040sh", "f1040es", "f1040x", "f1040v", "f1040sr", "f4868", "fw4", "fw9", "fw4p",
    "f1099msc", "f1099nec", "f1099int", "f1099div", "f1099r", "f1099k", "f1098", "f1098e", "f1098t",
    "f5498", "f8829", "f8863", "f2441", "f8962", "f8889", "f5695", "f8936", "f4562", "f8949", "f709",
    "f706", "f1041", "f1065", "f1120", "f1120s", "f941", "f940", "f8822", "f8888", "f9465", "f843",
    "f8379", "f8857", "fss4", "f2848", "f8821", "f4506t", "f8606", "f5329", "f3903", "f2106", "f8283",
    "f8332", "f8839", "f8880", "f8995", "f1040s8", "fw7", "f8862", "f1040sei", "f8615", "f4797",
    "f6251", "f8582", "f8812", "f4684", "f1310", "f56", "f8853",
    # instructions
    "i1040sc", "i8829", "i8962", "i8863", "i2441", "i8889", "i1040x", "i4562", "i8949", "i709",
    "i941", "i1065", "i9465", "i8606", "i5695", "i8283", "i4684",
    # publications
    "p17", "p501", "p503", "p505", "p523", "p525", "p526", "p527", "p530", "p544", "p550", "p551",
    "p554", "p555", "p559", "p560", "p570", "p575", "p587", "p590a", "p590b", "p596", "p915", "p936",
    "p946", "p969", "p970", "p463", "p334", "p583", "p15", "p519", "p54", "p1", "p594", "p556", "p3",
    "p502", "p504", "p521", "p535", "p547", "p915", "p4681", "p5", "p971", "p529", "p561", "p925",
]

ARXIV = {
    "gnn": "graph neural network",
    "qec": "quantum error correction",
    "exoplanet": "exoplanet atmosphere",
    "protein": "protein structure prediction",
    "asr": "speech recognition",
    "robot": "robotic manipulation reinforcement learning",
    "climate": "precipitation climate model",
    "darkmatter": "dark matter direct detection",
    "battery": "lithium battery electrolyte",
    "epidemic": "COVID-19 epidemic model",
    "gw": "gravitational wave black hole merger",
    "superres": "image super-resolution",
    "federated": "federated learning privacy",
    "mt": "low-resource machine translation",
    "superconductor": "superconductivity",
}
ARXIV_PER = 11

WIKI = {
    "dogs": ["Border Collie", "German Shepherd", "Labrador Retriever", "Dachshund", "Siberian Husky",
             "Beagle", "Australian Shepherd", "Great Dane", "Poodle", "Shetland Sheepdog", "Bulldog", "Whippet"],
    "volcanoes": ["Mount Vesuvius", "Krakatoa", "Mount St. Helens", "Eyjafjallajökull", "Mount Pinatubo",
                  "Mauna Loa", "Mount Etna", "Kīlauea", "Mount Fuji", "Popocatépetl", "Hekla", "Stromboli"],
    "pasta": ["Carbonara", "Lasagne", "Spaghetti alla puttanesca", "Pesto", "Ravioli", "Bolognese sauce",
              "Cacio e pepe", "Gnocchi", "Tortellini", "Fettuccine Alfredo", "Amatriciana", "Risotto"],
    "chess": ["Sicilian Defence", "Ruy Lopez", "Queen's Gambit", "French Defence", "Caro–Kann Defence",
              "King's Indian Defence", "Italian Game", "English Opening", "Nimzo-Indian Defence",
              "Scandinavian Defense", "King's Gambit", "Dutch Defence"],
    "emperors": ["Augustus", "Nero", "Trajan", "Hadrian", "Marcus Aurelius", "Caligula", "Constantine the Great",
                 "Diocletian", "Vespasian", "Tiberius", "Claudius", "Septimius Severus"],
    "languages": ["Python (programming language)", "Rust (programming language)", "C++", "JavaScript",
                  "Haskell", "Go (programming language)", "Fortran", "COBOL", "Swift (programming language)",
                  "Lisp (programming language)", "Kotlin (programming language)", "Erlang (programming language)"],
    "diseases": ["Malaria", "Tuberculosis", "Cholera", "Measles", "Influenza", "Dengue fever", "Smallpox",
                 "Typhoid fever", "Ebola", "Polio", "Yellow fever", "Leprosy"],
    "painters": ["Claude Monet", "Pierre-Auguste Renoir", "Edgar Degas", "Camille Pissarro", "Alfred Sisley",
                 "Berthe Morisot", "Mary Cassatt", "Gustave Caillebotte", "Édouard Manet", "Paul Cézanne",
                 "Frédéric Bazille", "Armand Guillaumin"],
    "cocktails": ["Martini (cocktail)", "Negroni", "Mojito", "Old fashioned (cocktail)", "Margarita",
                  "Daiquiri", "Manhattan (cocktail)", "Piña colada", "Cosmopolitan (cocktail)", "Whiskey sour",
                  "Gin and tonic", "Moscow mule"],
    "castles": ["Edinburgh Castle", "Windsor Castle", "Warwick Castle", "Caernarfon Castle", "Stirling Castle",
                "Bodiam Castle", "Dover Castle", "Conwy Castle", "Alnwick Castle", "Leeds Castle",
                "Harlech Castle", "Tower of London"],
    "probes": ["Voyager 1", "Voyager 2", "Cassini–Huygens", "New Horizons", "Galileo (spacecraft)", "Juno (spacecraft)",
               "Pioneer 10", "Rosetta (spacecraft)", "Messenger (spacecraft)", "Dawn (spacecraft)",
               "Magellan (spacecraft)", "Parker Solar Probe"],
    "instruments": ["Violin", "Cello", "Oboe", "Bassoon", "Trumpet", "French horn", "Harp", "Clarinet",
                    "Flute", "Double bass", "Trombone", "Viola"],
}
# Other languages: (lang, title, subject)
WIKI_INTL = [
    ("fr", "Vésuve", "volcanoes"), ("fr", "Etna", "volcanoes"), ("de", "Ätna", "volcanoes"),
    ("es", "Volcán Popocatépetl", "volcanoes"), ("de", "Hekla", "volcanoes"),
    ("fr", "Border collie", "dogs"), ("de", "Deutscher Schäferhund", "dogs"), ("es", "Labrador retriever", "dogs"),
    ("fr", "Teckel", "dogs"), ("de", "Siberian Husky", "dogs"),
    ("fr", "Paludisme", "diseases"), ("de", "Tuberkulose", "diseases"), ("es", "Cólera", "diseases"),
    ("fr", "Rougeole", "diseases"), ("es", "Dengue", "diseases"),
    ("fr", "Claude Monet", "painters"), ("de", "Edgar Degas", "painters"), ("es", "Pierre-Auguste Renoir", "painters"),
    ("fr", "Berthe Morisot", "painters"), ("fr", "Camille Pissarro", "painters"),
    ("fr", "Château d'Édimbourg", "castles"), ("de", "Windsor Castle", "castles"), ("es", "Torre de Londres", "castles"),
    ("fr", "Voyager 1", "probes"), ("de", "Cassini-Huygens", "probes"), ("es", "New Horizons", "probes"),
    ("fr", "Violon", "instruments"), ("de", "Violoncello", "instruments"), ("es", "Oboe", "instruments"),
    ("de", "Trompete", "instruments"),
]


def get(url, tries=4):
    for i in range(tries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA})
            with urllib.request.urlopen(req, timeout=60) as r:
                return r.read()
        except Exception as e:  # noqa: BLE001
            if i == tries - 1:
                print(f"  failed {url}: {e}", file=sys.stderr)
                return None
            time.sleep(3 * (i + 1))


def wiki_text(lang, title):
    q = urllib.parse.urlencode({"action": "query", "prop": "extracts", "explaintext": 1, "redirects": 1,
                                "titles": title, "format": "json", "formatversion": 2})
    time.sleep(1.5)  # Wikipedia answers 429 to a faster loop
    b = get(f"https://{lang}.wikipedia.org/w/api.php?{q}")
    if not b:
        return None
    pages = json.loads(b)["query"]["pages"]
    if not pages or "extract" not in pages[0] or len(pages[0]["extract"]) < 1500:
        return None
    return pages[0]["title"], pages[0]["extract"]


def write_docx(path, title, text):
    import docx
    d = docx.Document()
    d.add_heading(title, 0)
    for para in text.split("\n"):
        para = para.strip()
        if not para:
            continue
        m = re.match(r"^(=+)\s*(.*?)\s*=+$", para)
        if m:
            d.add_heading(m.group(2), min(len(m.group(1)), 4))
        else:
            d.add_paragraph(para)
    d.save(path)


def write_pdf(path, title, text):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_auto_page_break(True, 15)
    pdf.add_page()
    font = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"
    pdf.add_font("U", "", font)
    pdf.set_font("U", size=16)
    pdf.multi_cell(0, 9, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("U", size=11)
    for para in text.split("\n"):
        para = para.strip()
        if not para:
            continue
        m = re.match(r"^(=+)\s*(.*?)\s*=+$", para)
        if m:
            pdf.set_font("U", size=13)
            pdf.multi_cell(0, 7, m.group(2), new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("U", size=11)
        else:
            pdf.multi_cell(0, 5.5, para, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)
    pdf.output(path)


def main():
    out = sys.argv[1]
    docs_dir = os.path.join(out, "docs")
    os.makedirs(docs_dir, exist_ok=True)
    manifest_path = os.path.join(out, "corpus.json")
    corpus = json.load(open(manifest_path)) if os.path.exists(manifest_path) else {}

    def save():
        json.dump(corpus, open(manifest_path, "w"), indent=1, ensure_ascii=False)

    for name in dict.fromkeys(IRS):
        fn = f"irs-{name}.pdf"
        if fn in corpus:
            continue
        b = get(f"https://www.irs.gov/pub/irs-pdf/{name}.pdf")
        if b and b[:5] == b"%PDF-":
            open(os.path.join(docs_dir, fn), "wb").write(b)
            corpus[fn] = {"source": "irs", "subject": "irs", "title": name}
            save()
        time.sleep(0.3)
    print("irs", sum(1 for v in corpus.values() if v["source"] == "irs"))

    for subject, words in ARXIV.items():
        have = sum(1 for v in corpus.values() if v["source"] == "arxiv" and v["subject"] == subject)
        if have >= ARXIV_PER:
            continue
        q = " AND ".join(f'ti:"{w}"' if " " not in w else f"ti:{w}" for w in words.split())
        url = "http://export.arxiv.org/api/query?" + urllib.parse.urlencode(
            {"search_query": q, "max_results": ARXIV_PER + 4, "sortBy": "relevance"})
        b = get(url)
        time.sleep(3)
        if not b:
            continue
        ns = {"a": "http://www.w3.org/2005/Atom"}
        for e in ET.fromstring(b).findall("a:entry", ns):
            if have >= ARXIV_PER:
                break
            aid = e.find("a:id", ns).text.rsplit("/abs/", 1)[1]
            title = " ".join(e.find("a:title", ns).text.split())
            fn = "arxiv-" + aid.replace("/", "_") + ".pdf"
            if fn in corpus:
                continue
            pb = get(f"https://export.arxiv.org/pdf/{aid}")
            time.sleep(3)
            if pb and pb[:5] == b"%PDF-":
                open(os.path.join(docs_dir, fn), "wb").write(pb)
                corpus[fn] = {"source": "arxiv", "subject": subject, "title": title}
                have += 1
                save()
        print("arxiv", subject, have)

    n = 0
    for subject, titles in WIKI.items():
        for t in titles:
            n += 1
            ext = "docx" if n % 2 else "pdf"
            fn = "wiki-en-" + re.sub(r"[^A-Za-z0-9]+", "_", t).strip("_") + "." + ext
            if fn in corpus:
                continue
            r = wiki_text("en", t)
            if not r:
                print("  no text", t, file=sys.stderr)
                continue
            (write_docx if ext == "docx" else write_pdf)(os.path.join(docs_dir, fn), r[0], r[1])
            corpus[fn] = {"source": "wiki", "lang": "en", "subject": subject, "title": r[0]}
            save()
    for i, (lang, t, subject) in enumerate(WIKI_INTL):
        ext = "docx" if i % 2 else "pdf"
        fn = f"wiki-{lang}-" + re.sub(r"[^A-Za-z0-9]+", "_", t).strip("_") + "." + ext
        if fn in corpus:
            continue
        r = wiki_text(lang, t)
        if not r:
            print("  no text", lang, t, file=sys.stderr)
            continue
        (write_docx if ext == "docx" else write_pdf)(os.path.join(docs_dir, fn), r[0], r[1])
        corpus[fn] = {"source": "wiki", "lang": lang, "subject": subject, "title": r[0]}
        save()
    print("total", len(corpus))


if __name__ == "__main__":
    main()
